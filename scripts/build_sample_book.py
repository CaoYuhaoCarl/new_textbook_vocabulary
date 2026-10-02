#!/usr/bin/env python

import csv
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = ROOT / "outputs" / "019f8ca8-fe4a-7000-8d6b-d1ed952617fa"
SOURCE_CSV = OUTPUT_DIR / "vocabulary_classification_final.csv"
OUTPUT_DOCX = OUTPUT_DIR / "新教材同步词汇手册_样章.docx"

LATIN_FONT = "Calibri"
CJK_FONT = "Arial Unicode MS"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GOLD = "7A5A00"
GOLD_FILL = "FFF7DA"
WHITE = "FFFFFF"
BORDER = "D5DCE5"
BLACK = "1F2937"

FULL_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_TOP_BOTTOM = 80
CELL_MARGIN_START_END = 120


FOCUS_ENTRIES = {
    ("7A", "Starter Unit 1"): [
        {
            "word": "greet",
            "pattern": "greet sb; greet sb with a smile",
            "example": "We greeted each other at the school gate.",
            "translation": "我们在校门口互相问候。",
            "note": "greet是及物动词，后面可以直接接人。",
        },
        {
            "word": "each other",
            "pattern": "help / know / greet each other",
            "example": "The new classmates soon got to know each other.",
            "translation": "新同学们很快就互相认识了。",
            "note": "表示两人或多人之间的相互关系。",
        },
        {
            "word": "everyone",
            "pattern": "everyone + 单数谓语",
            "example": "Everyone is ready for the new term.",
            "translation": "每个人都为新学期做好了准备。",
            "note": "everyone作主语时，谓语动词通常使用单数形式。",
        },
        {
            "word": "conversation",
            "pattern": "have / start a conversation with sb",
            "example": "Ella started a conversation with the new student.",
            "translation": "埃拉和新同学聊了起来。",
            "note": "conversation是可数名词，注意搭配中的冠词a。",
        },
        {
            "word": "spell",
            "pattern": "spell a word / name",
            "example": "How do you spell your family name?",
            "translation": "你的姓怎么拼写？",
            "note": "询问拼写时常用How do you spell...?",
        },
        {
            "word": "start",
            "pattern": "start doing sth; start to do sth",
            "example": "The bell rang, and the class started.",
            "translation": "铃响了，课程开始了。",
            "note": "start后面可以接动名词，也可以接不定式。",
        },
    ],
    ("8B", "Unit 4"): [
        {
            "word": "wonder",
            "pattern": "wonder if / whether; natural wonder",
            "example": "I wonder how deep the lake is.",
            "translation": "我想知道这个湖有多深。",
            "note": "作动词时表示想知道，作名词时可表示奇观。",
        },
        {
            "word": "measurement",
            "pattern": "take measurements; measure the depth",
            "example": "Scientists took careful measurements of the cave.",
            "translation": "科学家对洞穴进行了仔细测量。",
            "note": "词族：measure v. 测量，measurement n. 测量结果。",
        },
        {
            "word": "below / above",
            "pattern": "below sea level; above the surface",
            "example": "The village lies 200 metres above sea level.",
            "translation": "这个村庄位于海拔200米处。",
            "note": "below和above都可以表示位置或数量上的低于和高于。",
        },
        {
            "word": "survive",
            "pattern": "survive the accident; survive in difficult conditions",
            "example": "Few plants can survive in such dry conditions.",
            "translation": "很少有植物能在如此干燥的环境中存活。",
            "note": "词族：survival n. 生存，survivor n. 幸存者。",
        },
        {
            "word": "determined",
            "pattern": "be determined to do sth",
            "example": "The climbers were determined to reach the top.",
            "translation": "登山者下定决心要到达山顶。",
            "note": "常与不定式连用，强调坚定的意愿。",
        },
        {
            "word": "risk",
            "pattern": "risk doing sth; at risk; take a risk",
            "example": "Do not risk swimming in dangerous water.",
            "translation": "不要冒险在危险水域游泳。",
            "note": "risk作动词时后接动名词，不接不定式。",
        },
        {
            "word": "attract",
            "pattern": "attract visitors; attract attention",
            "example": "The coral reef attracts travellers from around the world.",
            "translation": "这片珊瑚礁吸引着世界各地的旅行者。",
            "note": "词族：attraction n. 吸引力，attractive adj. 有吸引力的。",
        },
        {
            "word": "include",
            "pattern": "include sth; including + 名词",
            "example": "The price includes a guided tour of the coast.",
            "translation": "这个价格包括一次海岸导览。",
            "note": "include是谓语动词，including常用作介词。",
            "page_break_before": True,
        },
        {
            "word": "alive",
            "pattern": "stay alive; be alive",
            "example": "The turtle was still alive when people found it.",
            "translation": "人们发现这只海龟时，它还活着。",
            "note": "alive通常作表语，不直接放在名词前作定语。",
        },
    ],
}


EXERCISES = {
    ("7A", "Starter Unit 1"): [
        ("1", "Everyone ______ ready for the first class.  (is / are)"),
        ("2", "The two new students smiled and greeted ______ ______."),
        ("3", "How do you ______ the word conversation?"),
        ("4", "Complete the phrase: start a ______ with sb."),
        ("5", "Choose the correct word: We should help each ______.  (other / others)"),
        ("6", "Translate into English: 我们在校门口互相问候。"),
    ],
    ("8B", "Unit 4"): [
        ("1", "I ______ whether people can survive in such cold conditions."),
        ("2", "The lake is 50 metres ______ sea level.  (below / bottom)"),
        ("3", "The climbers were ______ to reach the top."),
        ("4", "Do not risk ______ alone in dangerous water.  (dive)"),
        ("5", "The park hopes to ______ more visitors next year."),
        ("6", "Complete the word family: deep adj. -> ______ n."),
        ("7", "Complete the word family: curious adj. -> ______ n."),
        ("8", "Complete the phrase: little by little = ______ by ______."),
        ("9", "Choose the correct word: The price includes / including lunch."),
        ("10", "Translate into English: 很少有植物能在这样的环境中存活。"),
    ],
}


ANSWERS = {
    ("7A", "Starter Unit 1"): [
        "1. is",
        "2. each other",
        "3. spell",
        "4. conversation",
        "5. other",
        "6. We greeted each other at the school gate.",
    ],
    ("8B", "Unit 4"): [
        "1. wonder",
        "2. below",
        "3. determined",
        "4. diving",
        "5. attract",
        "6. depth",
        "7. curiosity",
        "8. bit; bit",
        "9. includes",
        "10. Few plants can survive in such conditions.",
    ],
}


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_run_font(run, size=None, bold=None, italic=None, color=None, latin=LATIN_FONT, east_asia=CJK_FONT):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = rgb(color)


def set_style_font(style, size, color=BLACK, bold=None):
    style.font.name = LATIN_FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), LATIN_FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    if bold is not None:
        style.font.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=CELL_MARGIN_TOP_BOTTOM, start=CELL_MARGIN_START_END, bottom=CELL_MARGIN_TOP_BOTTOM, end=CELL_MARGIN_START_END):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size=4):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    if sum(widths_dxa) != FULL_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {FULL_WIDTH_DXA}: {widths_dxa}")

    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    if table.rows:
        set_repeat_table_header(table.rows[0])
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(FULL_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
    set_table_borders(table)


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_paragraph_border_bottom(paragraph, color=BORDER, size=8, space=6):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_page_number(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=9, color=MUTED)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def add_text(paragraph, text, size=10.5, bold=False, italic=False, color=BLACK):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return run


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    set_style_font(normal, 11, BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_1 = doc.styles["Heading 1"]
    set_style_font(heading_1, 16, BLUE, True)
    heading_1.paragraph_format.space_before = Pt(18)
    heading_1.paragraph_format.space_after = Pt(10)
    heading_1.paragraph_format.keep_with_next = True

    heading_2 = doc.styles["Heading 2"]
    set_style_font(heading_2, 13, BLUE, True)
    heading_2.paragraph_format.space_before = Pt(14)
    heading_2.paragraph_format.space_after = Pt(7)
    heading_2.paragraph_format.keep_with_next = True

    heading_3 = doc.styles["Heading 3"]
    set_style_font(heading_3, 12, DARK_BLUE, True)
    heading_3.paragraph_format.space_before = Pt(10)
    heading_3.paragraph_format.space_after = Pt(5)
    heading_3.paragraph_format.keep_with_next = True

    for doc_section in doc.sections:
        header_p = doc_section.header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header_p.paragraph_format.space_after = Pt(0)
        add_text(header_p, "新教材同步词汇手册  |  样章", size=8.5, color=MUTED)

        footer_p = doc_section.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer_p.paragraph_format.space_before = Pt(0)
        footer_p.paragraph_format.space_after = Pt(0)
        add_text(footer_p, "样章  |  ", size=9, color=MUTED)
        add_page_number(footer_p)


def add_cover(doc):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(88)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    add_text(kicker, "VOCABULARY SAMPLE", size=10, bold=True, color=GOLD)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    add_text(title, "新教材同步词汇手册", size=28, bold=True, color=INK)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(6)
    add_text(subtitle, "教材同步  |  浙江中考导向  |  分级记忆", size=13, color=DARK_BLUE)

    sample = doc.add_paragraph()
    sample.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sample.paragraph_format.space_after = Pt(42)
    add_text(sample, "样章 01  ·  7A Starter Unit 1 + 8B Unit 4", size=11, bold=True, color=MUTED)

    statement = doc.add_paragraph()
    statement.alignment = WD_ALIGN_PARAGRAPH.CENTER
    statement.paragraph_format.space_after = Pt(70)
    statement.paragraph_format.line_spacing = 1.35
    add_text(statement, "从教材词表出发，先剔除不要求背诵的专名，\n再把真正需要学习的词汇组织成可复习、可检测的单元。", size=12, color=BLACK)

    edition = doc.add_paragraph()
    edition.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(edition, "内部样章  |  2026年7月", size=10, color=MUTED)
    doc.add_page_break()


def add_heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def add_lead_callout(doc, label, text, fill=CALLOUT, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [FULL_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    add_text(paragraph, f"{label}  ", size=10, bold=True, color=accent)
    add_text(paragraph, text, size=10, color=BLACK)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_how_to_use(doc):
    add_heading(doc, "如何使用这本词汇手册", 1)
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(10)
    add_text(intro, "本书保留教材顺序，同时把学习要求分成两个层级。专名不进入学生背诵表。", size=11)

    table = doc.add_table(rows=3, cols=3)
    set_table_geometry(table, [1200, 2100, 6060])
    headers = ["层级", "学习要求", "使用方式"]
    values = [
        ["A核心", "会读、会写、会用", "掌握基本词义、常用搭配和语境使用。"],
        ["B认读", "会读、会认、理解", "理解课文和阅读语境，不要求重点默写。"],
    ]
    for index, value in enumerate(headers):
        set_cell_text(table.cell(0, index), value, bold=True, color=WHITE, align="center")
        set_cell_shading(table.cell(0, index), BLUE)
    for row_index, row_values in enumerate(values, start=1):
        for col_index, value in enumerate(row_values):
            fill = LIGHT_BLUE if row_index == 1 else GOLD_FILL
            set_cell_shading(table.cell(row_index, col_index), fill)
            set_cell_text(
                table.cell(row_index, col_index),
                value,
                bold=col_index == 0,
                color=DARK_BLUE if row_index == 1 else GOLD,
                align="center" if col_index < 2 else "left",
            )

    add_heading(doc, "六次复习法", 2)
    tracker = doc.add_table(rows=2, cols=7)
    set_table_geometry(tracker, [900, 1410, 1410, 1410, 1410, 1410, 1410])
    labels = ["复习", "第1天", "第2天", "第4天", "第7天", "第14天", "第30天"]
    checks = ["完成", "□", "□", "□", "□", "□", "□"]
    for i, label in enumerate(labels):
        set_cell_shading(tracker.cell(0, i), LIGHT_BLUE)
        set_cell_text(tracker.cell(0, i), label, bold=True, color=DARK_BLUE, align="center")
        set_cell_text(tracker.cell(1, i), checks[i], size=11, align="center")

    add_lead_callout(doc, "筛词说明", "人名、城市、景点、作品名等专名已从主词表中移除。常见国家、大洲、节日和赛事保留在B类认读区。")
    doc.add_page_break()


def set_cell_text(cell, text, size=9.5, bold=False, italic=False, color=BLACK, align="left"):
    paragraph = cell.paragraphs[0]
    paragraph.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    add_text(paragraph, str(text or ""), size=size, bold=bold, italic=italic, color=color)


def add_unit_overview(doc, book, unit, rows, theme, suggested_days, page_break_before=False):
    retained = [row for row in rows if row["classification"] != "C剔除"]
    core = [row for row in rows if row["classification"] == "A核心"]
    recognition = [row for row in rows if row["classification"] == "B认读"]
    excluded = [row for row in rows if row["classification"] == "C剔除"]

    kicker = doc.add_paragraph()
    kicker.paragraph_format.page_break_before = page_break_before
    kicker.paragraph_format.space_after = Pt(2)
    add_text(kicker, f"{book}  |  {unit}", size=10, bold=True, color=GOLD)

    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(8)
    add_text(heading, theme, size=22, bold=True, color=INK)
    set_paragraph_border_bottom(heading, color=LIGHT_BLUE, size=12, space=6)

    summary_table = doc.add_table(rows=2, cols=5)
    set_table_geometry(summary_table, [1700, 1700, 1700, 1700, 2560])
    summary_headers = ["教材词条", "A核心", "B认读", "已剔除", "建议节奏"]
    summary_values = [str(len(rows)), str(len(core)), str(len(recognition)), str(len(excluded)), suggested_days]
    for i, value in enumerate(summary_headers):
        set_cell_shading(summary_table.cell(0, i), LIGHT_BLUE)
        set_cell_text(summary_table.cell(0, i), value, bold=True, color=DARK_BLUE, align="center")
        set_cell_text(summary_table.cell(1, i), summary_values[i], size=11, bold=True, color=INK, align="center")

    add_lead_callout(
        doc,
        "学习目标",
        f"本单元保留{len(retained)}条学习词目。先掌握A类核心，再完成B类认读，最后用语境练习检查。",
    )

    if excluded:
        examples = "、".join(row["word"] for row in excluded[:5])
        add_lead_callout(
            doc,
            "本单元筛选",
            f"已隐藏{len(excluded)}条不要求背诵的专名，例如{examples}。",
            fill=LIGHT_GRAY,
            accent=MUTED,
        )


def add_word_table(doc, title, rows, level):
    add_heading(doc, title, 2)
    if not rows:
        paragraph = doc.add_paragraph()
        add_text(paragraph, "本单元无此类词目。", color=MUTED)
        return

    table = doc.add_table(rows=1, cols=5)
    widths = [2000, 1750, 680, 4090, 840]
    headers = ["词汇", "音标", "词性", "教材释义", "页码"]
    for i, header in enumerate(headers):
        set_cell_text(table.cell(0, i), header, bold=True, color=WHITE, align="center")
        set_cell_shading(table.cell(0, i), BLUE if level == "A核心" else GOLD)
    set_repeat_table_header(table.rows[0])

    for row_index, item in enumerate(rows, start=1):
        cells = table.add_row().cells
        fill = WHITE if row_index % 2 else ("F7FAFC" if level == "A核心" else "FFFBEB")
        for cell in cells:
            set_cell_shading(cell, fill)
        part_of_speech = item["part_of_speech"]
        definition = item["chinese_definition"]
        variant_match = re.match(r"^\(\s*=\s*([^)]+)\)\s*(.*)$", part_of_speech)
        if variant_match:
            variant = variant_match.group(1).strip()
            part_of_speech = variant_match.group(2).strip()
            definition = f"{definition}；亦作 {variant}"
        set_cell_text(cells[0], item["word"], bold=True, color=DARK_BLUE if level == "A核心" else GOLD)
        set_cell_text(cells[1], item["pronunciation"], size=9)
        set_cell_text(cells[2], part_of_speech, size=8.8, align="center")
        set_cell_text(cells[3], definition, size=9.3)
        set_cell_text(cells[4], item["page_number"], size=9, align="center")
    set_table_geometry(table, widths)


def add_focus_entries(doc, book, unit):
    add_heading(doc, "重点用法", 2)
    for item in FOCUS_ENTRIES[(book, unit)]:
        title = doc.add_paragraph()
        title.paragraph_format.page_break_before = item.get("page_break_before", False)
        title.paragraph_format.space_before = Pt(7)
        title.paragraph_format.space_after = Pt(3)
        title.paragraph_format.keep_with_next = True
        set_paragraph_shading(title, LIGHT_BLUE)
        add_text(title, item["word"], size=12, bold=True, color=DARK_BLUE)

        detail = doc.add_paragraph()
        detail.paragraph_format.keep_with_next = True
        detail.paragraph_format.space_after = Pt(2)
        add_text(detail, "搭配  ", size=9.5, bold=True, color=GOLD)
        add_text(detail, item["pattern"], size=9.5)

        example = doc.add_paragraph()
        example.paragraph_format.keep_with_next = True
        example.paragraph_format.space_after = Pt(1)
        add_text(example, item["example"], size=10.2, bold=True, color=INK)

        translation = doc.add_paragraph()
        translation.paragraph_format.keep_with_next = True
        translation.paragraph_format.space_after = Pt(1)
        add_text(translation, item["translation"], size=9.5, color=MUTED)

        note = doc.add_paragraph()
        note.paragraph_format.space_after = Pt(4)
        add_text(note, "提示  ", size=9.5, bold=True, color=BLUE)
        add_text(note, item["note"], size=9.5)


def add_exercises(doc, book, unit):
    add_heading(doc, "单元检测", 2)
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(6)
    add_text(intro, "先独立完成，再到书末核对答案。", size=9.5, color=MUTED)

    items = EXERCISES[(book, unit)]
    table = doc.add_table(rows=len(items), cols=2)
    set_table_geometry(table, [600, 8760])
    for index, (number, prompt) in enumerate(items):
        fill = LIGHT_GRAY if index % 2 == 0 else WHITE
        set_cell_shading(table.cell(index, 0), fill)
        set_cell_shading(table.cell(index, 1), fill)
        set_cell_text(table.cell(index, 0), number, bold=True, color=BLUE, align="center")
        set_cell_text(table.cell(index, 1), prompt, size=10)

    add_heading(doc, "复习记录", 3)
    tracker = doc.add_table(rows=2, cols=7)
    set_table_geometry(tracker, [900, 1410, 1410, 1410, 1410, 1410, 1410])
    labels = ["复习", "第1天", "第2天", "第4天", "第7天", "第14天", "第30天"]
    for i, label in enumerate(labels):
        set_cell_shading(tracker.cell(0, i), LIGHT_BLUE)
        set_cell_text(tracker.cell(0, i), label, bold=True, color=DARK_BLUE, align="center")
        set_cell_text(tracker.cell(1, i), "完成" if i == 0 else "□", size=11, align="center")


def add_unit(doc, book, unit, rows, theme, suggested_days, page_break_before=False):
    add_unit_overview(doc, book, unit, rows, theme, suggested_days, page_break_before)
    core = [row for row in rows if row["classification"] == "A核心"]
    recognition = [row for row in rows if row["classification"] == "B认读"]
    add_word_table(doc, "A类核心词", core, "A核心")
    add_word_table(doc, "B类认读词", recognition, "B认读")
    add_focus_entries(doc, book, unit)
    add_exercises(doc, book, unit)


def add_answer_key(doc):
    heading = add_heading(doc, "参考答案", 1)
    heading.paragraph_format.page_break_before = True
    for book, unit in (("7A", "Starter Unit 1"), ("8B", "Unit 4")):
        add_heading(doc, f"{book} {unit}", 2)
        answers = ANSWERS[(book, unit)]
        table = doc.add_table(rows=len(answers), cols=1)
        set_table_geometry(table, [FULL_WIDTH_DXA])
        for index, answer in enumerate(answers):
            set_cell_shading(table.cell(index, 0), LIGHT_GRAY if index % 2 == 0 else WHITE)
            set_cell_text(table.cell(index, 0), answer, size=10)

    add_lead_callout(doc, "样章说明", "本样章用于确认筛词原则、内容密度、词条层级与练习形式。全书将在样章通过后按同一结构生成。")


def read_rows():
    with SOURCE_CSV.open("r", encoding="utf-8-sig", newline="") as file:
        return list(csv.DictReader(file))


def build_document():
    all_rows = read_rows()
    selected = {}
    for book, unit in (("7A", "Starter Unit 1"), ("8B", "Unit 4")):
        selected[(book, unit)] = [
            row for row in all_rows if row["book"] == book and row["unit"] == unit
        ]
        if not selected[(book, unit)]:
            raise RuntimeError(f"No vocabulary rows found for {book} {unit}")

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "新教材同步词汇手册样章"
    doc.core_properties.subject = "7A Starter Unit 1 与 8B Unit 4"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""

    add_cover(doc)
    add_how_to_use(doc)
    add_unit(
        doc,
        "7A",
        "Starter Unit 1",
        selected[("7A", "Starter Unit 1")],
        "打招呼与开始交流",
        "2天学习 + 6次复习",
    )
    add_unit(
        doc,
        "8B",
        "Unit 4",
        selected[("8B", "Unit 4")],
        "地理奇观与探索",
        "5天学习 + 6次复习",
        page_break_before=True,
    )
    add_answer_key(doc)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    print(f"output={OUTPUT_DOCX}")
    print(f"7A_rows={len(selected[('7A', 'Starter Unit 1')])}")
    print(f"8B_rows={len(selected[('8B', 'Unit 4')])}")


if __name__ == "__main__":
    build_document()
