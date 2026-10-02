#!/usr/bin/env python

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = ROOT / "outputs" / "019f8ca8-fe4a-7000-8d6b-d1ed952617fa"
SOURCE_DOCX = OUTPUT_DIR / "新教材同步词汇手册_完整本.docx"
OUTPUT_DOCX = OUTPUT_DIR / "新教材同步词汇手册_中文释义默写版.docx"

WORD_TABLE_HEADERS = ["词汇", "音标", "词性", "教材释义", "页码"]
ANSWER_TABLE_HEADERS = ["单元", "核心词自测答案"]
TEST_TABLE_HEADERS = ["序号", "拼写自测"]
EXPECTED_WORD_TABLES = 94
EXPECTED_ENTRIES = 2445
EXPECTED_PRIORITY_CALLOUTS = 47
EXPECTED_ANSWER_ROWS = 47

POS_TOKEN_PATTERN = re.compile(
    r"modal\s+v\.|linking\s+v\.|aux(?:iliary)?\.?\s+v\.|"
    r"phrasal\s+v\.|interj\.|pron\.|prep\.|conj\.|adj\.|adv\.|"
    r"num\.|art\.|n\.|v\.",
    re.IGNORECASE,
)
POS_CANONICAL = {
    "modal v.": "modal v.",
    "linking v.": "linking v.",
    "aux v.": "aux. v.",
    "aux. v.": "aux. v.",
    "auxiliary v.": "aux. v.",
    "phrasal v.": "phrasal v.",
    "interj.": "interj.",
    "pron.": "pron.",
    "prep.": "prep.",
    "conj.": "conj.",
    "adj.": "adj.",
    "adv.": "adv.",
    "num.": "num.",
    "art.": "art.",
    "n.": "n.",
    "v.": "v.",
}


def text_nodes(element):
    return list(element.iter(qn("w:t")))


def clear_visible_text(element):
    for node in text_nodes(element):
        node.text = ""


def replace_visible_text(element, replacement):
    nodes = text_nodes(element)
    if not nodes:
        if not replacement:
            return
        raise RuntimeError("Cannot replace text because no text node exists.")
    nodes[0].text = replacement
    for node in nodes[1:]:
        node.text = ""


def sanitize_part_of_speech(value):
    tokens = []
    for match in POS_TOKEN_PATTERN.finditer(value or ""):
        raw = re.sub(r"\s+", " ", match.group(0).lower()).strip()
        canonical = POS_CANONICAL.get(raw, raw)
        if canonical not in tokens:
            tokens.append(canonical)
    return " & ".join(tokens)


def remove_latin_parentheticals(value):
    text = value
    patterns = [
        re.compile(r"\([^()]*[A-Za-z][^()]*\)"),
        re.compile(r"（[^（）]*[A-Za-z][^（）]*）"),
    ]
    changed = True
    while changed:
        changed = False
        for pattern in patterns:
            updated = pattern.sub("", text)
            if updated != text:
                text = updated
                changed = True
    return text


def sanitize_chinese_definition(value):
    text = value or ""
    text = re.sub(
        r"[；，,:：]?\s*亦作\s*[A-Za-z][A-Za-z\s'’/-]*",
        "",
        text,
    )
    text = remove_latin_parentheticals(text)
    text = re.sub(r"/[^/\n]+/", "", text)
    text = re.sub(r"[A-Za-z]+(?:[-'’][A-Za-z]+)*", "", text)
    text = re.sub(r"[=]+", "", text)
    text = re.sub(r"\(\s*[.&/;,，；:：-]*\s*\)", "", text)
    text = re.sub(r"（\s*[.&/;,，；:：-]*\s*）", "", text)
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\s+([，；。,:：])", r"\1", text)
    text = re.sub(r"([，；,:：])\s*([，；,:：])+", r"\1", text)
    text = re.sub(r"^[\s.&/;,，；:：-]+", "", text)
    text = re.sub(r"[\s.&/;,，；:：-]+$", "", text)
    return text.strip()


def sanitize_test_prompt(value):
    body = re.sub(r"^_+\s*", "", value or "")
    part_of_speech = sanitize_part_of_speech(body)
    chinese_definition = sanitize_chinese_definition(body)
    pieces = [
        piece for piece in (part_of_speech, chinese_definition) if piece
    ]
    return "________________  " + " ".join(pieces)


def table_headers(table):
    return [cell.text.strip() for cell in table.rows[0].cells]


def transform_document():
    if not SOURCE_DOCX.exists():
        raise RuntimeError(f"Source document not found: {SOURCE_DOCX}")

    doc = Document(SOURCE_DOCX)
    word_table_count = 0
    entry_count = 0
    priority_count = 0
    answer_row_count = 0
    test_row_count = 0

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == "VOCABULARY HANDBOOK":
            replace_visible_text(paragraph._p, "中文释义默写版")
        elif text == "完整本  |  2026年7月":
            replace_visible_text(paragraph._p, "中文释义默写版  |  2026年7月")

    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            if "新教材同步词汇手册" in paragraph.text:
                replace_visible_text(
                    paragraph._p,
                    "新教材同步词汇手册  |  中文释义默写版",
                )
        for paragraph in section.footer.paragraphs:
            if "完整本" in paragraph.text:
                nodes = text_nodes(paragraph._p)
                if nodes:
                    nodes[0].text = "默写版  |  "

    for table in doc.tables:
        headers = table_headers(table)
        if headers == WORD_TABLE_HEADERS:
            word_table_count += 1
            for row in table.rows[1:]:
                clear_visible_text(row.cells[0]._tc)
                clear_visible_text(row.cells[1]._tc)
                replace_visible_text(
                    row.cells[2]._tc,
                    sanitize_part_of_speech(row.cells[2].text),
                )
                replace_visible_text(
                    row.cells[3]._tc,
                    sanitize_chinese_definition(row.cells[3].text),
                )
                entry_count += 1
            continue

        if headers == TEST_TABLE_HEADERS:
            for row in table.rows[1:]:
                replace_visible_text(
                    row.cells[1]._tc,
                    sanitize_test_prompt(row.cells[1].text),
                )
                test_row_count += 1
            continue

        if headers == ANSWER_TABLE_HEADERS:
            for row in table.rows[1:]:
                clear_visible_text(row.cells[1]._tc)
                answer_row_count += 1
            continue

        if len(table.rows) == 1 and len(table.columns) == 1:
            cell = table.cell(0, 0)
            if cell.text.strip().startswith("优先记忆"):
                nodes = text_nodes(cell._tc)
                if not nodes:
                    raise RuntimeError("Priority callout has no text node.")
                nodes[0].text = "优先记忆  "
                for node in nodes[1:]:
                    node.text = ""
                priority_count += 1

    if word_table_count != EXPECTED_WORD_TABLES:
        raise RuntimeError(
            f"Expected {EXPECTED_WORD_TABLES} word tables, found {word_table_count}"
        )
    if entry_count != EXPECTED_ENTRIES:
        raise RuntimeError(
            f"Expected {EXPECTED_ENTRIES} entries, found {entry_count}"
        )
    if priority_count != EXPECTED_PRIORITY_CALLOUTS:
        raise RuntimeError(
            f"Expected {EXPECTED_PRIORITY_CALLOUTS} priority callouts, "
            f"found {priority_count}"
        )
    if answer_row_count != EXPECTED_ANSWER_ROWS:
        raise RuntimeError(
            f"Expected {EXPECTED_ANSWER_ROWS} answer rows, found {answer_row_count}"
        )

    doc.core_properties.title = "新教材同步词汇手册中文释义默写版"
    doc.core_properties.subject = "保留词性和中文释义的全册默写版"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)

    print(f"output={OUTPUT_DOCX}")
    print(f"word_tables={word_table_count}")
    print(f"entries_cleared={entry_count}")
    print(f"priority_callouts_cleared={priority_count}")
    print(f"answer_rows_cleared={answer_row_count}")
    print(f"test_prompts_sanitized={test_row_count}")


if __name__ == "__main__":
    transform_document()
