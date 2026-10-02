#!/usr/bin/env python

import argparse
import csv
import math
import re
from collections import OrderedDict
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

import build_sample_book as base


ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = ROOT / "outputs" / "019f8ca8-fe4a-7000-8d6b-d1ed952617fa"
SOURCE_CSV = OUTPUT_DIR / "vocabulary_classification_final.csv"
OUTPUT_DOCX = OUTPUT_DIR / "新教材同步词汇手册_完整本.docx"

BOOK_ORDER = ["7A", "7B", "8A", "8B", "9A", "9B"]
BOOK_NAMES = {
    "7A": "七年级上册",
    "7B": "七年级下册",
    "8A": "八年级上册",
    "8B": "八年级下册",
    "9A": "九年级上册",
    "9B": "九年级下册",
}

EXPECTED_TOTAL = 2681
EXPECTED_RETAINED = 2445
EXPECTED_EXCLUDED = 236
EXPECTED_CORE = 1797
EXPECTED_RECOGNITION = 648
EXPECTED_UNITS = 47


def parse_args():
    parser = argparse.ArgumentParser(description="Build the complete vocabulary handbook.")
    parser.add_argument(
        "--output",
        type=Path,
        default=OUTPUT_DOCX,
        help="Output DOCX path.",
    )
    return parser.parse_args()


def clean_text(value):
    return re.sub(r"\s+", " ", value or "").strip()


def unit_sort_key(unit):
    starter_match = re.fullmatch(r"Starter Unit\s+(\d+)", unit)
    if starter_match:
        return (0, int(starter_match.group(1)))
    unit_match = re.fullmatch(r"Unit\s+(\d+)", unit)
    if unit_match:
        return (1, int(unit_match.group(1)))
    return (2, unit)


def rank_value(row):
    value = clean_text(row.get("exam_frequency_rank"))
    return int(value) if value.isdigit() else 10**9


def study_days(rows):
    retained_count = sum(
        row["classification"] in {"A核心", "B认读"} for row in rows
    )
    return max(1, math.ceil(retained_count / 15))


def read_rows():
    with SOURCE_CSV.open("r", encoding="utf-8-sig", newline="") as file:
        rows = list(csv.DictReader(file))
    if len(rows) != EXPECTED_TOTAL:
        raise RuntimeError(f"Expected {EXPECTED_TOTAL} rows, found {len(rows)}")
    return rows


def group_rows(rows):
    grouped = {}
    for row in rows:
        key = (row["book"], row["unit"])
        grouped.setdefault(key, []).append(row)

    ordered = OrderedDict()
    for book in BOOK_ORDER:
        units = sorted(
            (unit for grouped_book, unit in grouped if grouped_book == book),
            key=unit_sort_key,
        )
        ordered[book] = OrderedDict(
            (unit, grouped[(book, unit)]) for unit in units
        )
    return ordered


def validate_rows(grouped):
    all_rows = [
        row
        for units in grouped.values()
        for rows in units.values()
        for row in rows
    ]
    counts = {
        "A核心": sum(row["classification"] == "A核心" for row in all_rows),
        "B认读": sum(row["classification"] == "B认读" for row in all_rows),
        "C剔除": sum(row["classification"] == "C剔除" for row in all_rows),
    }
    retained = counts["A核心"] + counts["B认读"]
    unit_count = sum(len(units) for units in grouped.values())

    expected = {
        "A核心": EXPECTED_CORE,
        "B认读": EXPECTED_RECOGNITION,
        "C剔除": EXPECTED_EXCLUDED,
    }
    if counts != expected:
        raise RuntimeError(f"Unexpected class counts: {counts}")
    if retained != EXPECTED_RETAINED:
        raise RuntimeError(f"Unexpected retained count: {retained}")
    if unit_count != EXPECTED_UNITS:
        raise RuntimeError(f"Unexpected unit count: {unit_count}")
    if set(grouped) != set(BOOK_ORDER):
        raise RuntimeError(f"Unexpected book set: {list(grouped)}")


def clear_paragraph_content(paragraph):
    paragraph_element = paragraph._p
    for child in list(paragraph_element):
        if child.tag != qn("w:pPr"):
            paragraph_element.remove(child)


def set_update_fields(doc):
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def configure_full_document(doc):
    base.configure_document(doc)
    set_update_fields(doc)
    for section in doc.sections:
        header = section.header.paragraphs[0]
        clear_paragraph_content(header)
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header.paragraph_format.space_after = Pt(0)
        base.add_text(
            header,
            "新教材同步词汇手册  |  完整本",
            size=8.5,
            color=base.MUTED,
        )

        footer = section.footer.paragraphs[0]
        clear_paragraph_content(footer)
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.paragraph_format.space_before = Pt(0)
        footer.paragraph_format.space_after = Pt(0)
        base.add_text(footer, "完整本  |  ", size=9, color=base.MUTED)
        base.add_page_number(footer)


def add_cover(doc):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(82)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    base.add_text(
        kicker,
        "VOCABULARY HANDBOOK",
        size=10,
        bold=True,
        color=base.GOLD,
    )

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    base.add_text(
        title,
        "新教材同步词汇手册",
        size=28,
        bold=True,
        color=base.INK,
    )

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    base.add_text(
        subtitle,
        "教材同步  |  浙江中考导向  |  分级记忆",
        size=13,
        color=base.DARK_BLUE,
    )

    scope = doc.add_paragraph()
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scope.paragraph_format.space_after = Pt(42)
    base.add_text(
        scope,
        "覆盖 7A - 9B  |  47个单元  |  2,445条学习词汇",
        size=11,
        bold=True,
        color=base.MUTED,
    )

    statement = doc.add_paragraph()
    statement.alignment = WD_ALIGN_PARAGRAPH.CENTER
    statement.paragraph_format.space_after = Pt(70)
    statement.paragraph_format.line_spacing = 1.35
    base.add_text(
        statement,
        "从教材词表出发，剔除不要求背诵的专名，\n"
        "按核心掌握与认读理解组织每个单元。",
        size=12,
        color=base.BLACK,
    )

    edition = doc.add_paragraph()
    edition.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.add_text(
        edition,
        "完整本  |  2026年7月",
        size=10,
        color=base.MUTED,
    )
    doc.add_page_break()


def set_header_row(table, headers, fill=base.BLUE):
    for index, value in enumerate(headers):
        base.set_cell_shading(table.cell(0, index), fill)
        base.set_cell_text(
            table.cell(0, index),
            value,
            bold=True,
            color=base.WHITE,
            align="center",
        )


def add_how_to_use(doc):
    base.add_how_to_use(doc)


def add_book_summary_table(doc, grouped):
    table = doc.add_table(rows=1, cols=5)
    base.set_table_geometry(table, [1800, 1500, 2000, 2000, 2060])
    set_header_row(table, ["册次", "单元数", "学习词条", "A类核心", "B类认读"])

    for row_index, book in enumerate(BOOK_ORDER, start=1):
        rows = [
            row
            for unit_rows in grouped[book].values()
            for row in unit_rows
        ]
        retained = [
            row for row in rows if row["classification"] in {"A核心", "B认读"}
        ]
        values = [
            f"{BOOK_NAMES[book]} {book}",
            str(len(grouped[book])),
            str(len(retained)),
            str(sum(row["classification"] == "A核心" for row in rows)),
            str(sum(row["classification"] == "B认读" for row in rows)),
        ]
        cells = table.add_row().cells
        fill = base.WHITE if row_index % 2 else base.LIGHT_GRAY
        for column_index, value in enumerate(values):
            base.set_cell_shading(cells[column_index], fill)
            base.set_cell_text(
                cells[column_index],
                value,
                size=9.5,
                bold=column_index == 0,
                color=base.DARK_BLUE if column_index == 0 else base.BLACK,
                align="center",
            )
    base.set_table_geometry(table, [1800, 1500, 2000, 2000, 2060])


def add_contents_table(doc, grouped):
    base.add_heading(doc, "全书目录", 1)
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(8)
    base.add_text(
        intro,
        "可通过Word导航窗格中的标题快速定位册次和单元。",
        size=10,
        color=base.MUTED,
    )

    table = doc.add_table(rows=1, cols=4)
    widths = [1200, 3000, 1000, 4160]
    base.set_table_geometry(table, widths)
    set_header_row(table, ["册次", "单元", "词条", "学习节奏"])

    row_index = 0
    for book in BOOK_ORDER:
        for unit, rows in grouped[book].items():
            row_index += 1
            retained = [
                row
                for row in rows
                if row["classification"] in {"A核心", "B认读"}
            ]
            values = [
                book,
                unit,
                str(len(retained)),
                f"{study_days(rows)}天学习 + 6次复习",
            ]
            cells = table.add_row().cells
            fill = base.WHITE if row_index % 2 else base.LIGHT_GRAY
            for column_index, value in enumerate(values):
                base.set_cell_shading(cells[column_index], fill)
                base.set_cell_text(
                    cells[column_index],
                    value,
                    size=9.2,
                    bold=column_index in {0, 1},
                    color=base.DARK_BLUE if column_index in {0, 1} else base.BLACK,
                    align="center" if column_index != 3 else "left",
                )
    base.set_table_geometry(table, widths)


def add_front_matter(doc, grouped):
    base.add_heading(doc, "全书导览", 1)
    base.add_lead_callout(
        doc,
        "筛词结果",
        "原始词表共2,681条。完整本保留2,445条学习词汇，剔除236条人名、"
        "具体地名、作品名和低考试价值文化专名。",
    )
    add_book_summary_table(doc, grouped)
    base.add_lead_callout(
        doc,
        "学习原则",
        "A类核心词要求会读、会写、会用。B类认读词要求在课文和阅读语境中能够识别与理解。",
        fill=base.GOLD_FILL,
        accent=base.GOLD,
    )
    add_contents_table(doc, grouped)


def add_book_divider(doc, book, units):
    doc.add_page_break()

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(62)
    kicker.paragraph_format.space_after = Pt(8)
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.add_text(
        kicker,
        f"BOOK {book}",
        size=10,
        bold=True,
        color=base.GOLD,
    )

    title = doc.add_paragraph(style="Heading 1")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    base.add_text(
        title,
        f"{BOOK_NAMES[book]}  {book}",
        size=26,
        bold=True,
        color=base.INK,
    )

    all_rows = [row for rows in units.values() for row in rows]
    retained = [
        row for row in all_rows if row["classification"] in {"A核心", "B认读"}
    ]
    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.paragraph_format.space_after = Pt(26)
    base.add_text(
        summary,
        f"{len(units)}个单元  |  {len(retained)}条学习词汇  |  "
        f"A类{sum(row['classification'] == 'A核心' for row in all_rows)}条  |  "
        f"B类{sum(row['classification'] == 'B认读' for row in all_rows)}条",
        size=10.5,
        color=base.MUTED,
    )

    table = doc.add_table(rows=1, cols=5)
    widths = [2900, 1300, 1300, 1300, 2560]
    base.set_table_geometry(table, widths)
    set_header_row(table, ["单元", "词条", "A类", "B类", "建议节奏"])

    for row_index, (unit, rows) in enumerate(units.items(), start=1):
        core_count = sum(row["classification"] == "A核心" for row in rows)
        recognition_count = sum(row["classification"] == "B认读" for row in rows)
        values = [
            unit,
            str(core_count + recognition_count),
            str(core_count),
            str(recognition_count),
            f"{study_days(rows)}天学习",
        ]
        cells = table.add_row().cells
        fill = base.WHITE if row_index % 2 else base.LIGHT_GRAY
        for column_index, value in enumerate(values):
            base.set_cell_shading(cells[column_index], fill)
            base.set_cell_text(
                cells[column_index],
                value,
                size=9.5,
                bold=column_index == 0,
                color=base.DARK_BLUE if column_index == 0 else base.BLACK,
                align="center",
            )
    base.set_table_geometry(table, widths)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(16)
    base.add_text(
        note,
        "建议先完成A类核心词，再阅读B类认读词，并按复习记录进行循环巩固。",
        size=10,
        color=base.MUTED,
    )
    doc.add_page_break()


def add_unit_overview(doc, book, unit, rows, page_break_before):
    kicker = doc.add_paragraph()
    kicker.paragraph_format.page_break_before = page_break_before
    kicker.paragraph_format.space_after = Pt(2)
    base.add_text(
        kicker,
        f"{book}  |  {unit}",
        size=10,
        bold=True,
        color=base.GOLD,
    )

    heading = doc.add_paragraph(f"{book} {unit}", style="Heading 1")
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(3)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(8)
    base.add_text(
        subtitle,
        "核心词汇与认读词",
        size=19,
        bold=True,
        color=base.INK,
    )
    base.set_paragraph_border_bottom(
        subtitle,
        color=base.LIGHT_BLUE,
        size=12,
        space=6,
    )

    core = [row for row in rows if row["classification"] == "A核心"]
    recognition = [row for row in rows if row["classification"] == "B认读"]
    excluded = [row for row in rows if row["classification"] == "C剔除"]
    retained = core + recognition
    pace = f"{study_days(rows)}天学习 + 6次复习"

    summary_table = doc.add_table(rows=2, cols=5)
    widths = [1700, 1700, 1700, 1700, 2560]
    base.set_table_geometry(summary_table, widths)
    headers = ["教材词条", "A核心", "B认读", "已剔除", "建议节奏"]
    values = [
        str(len(rows)),
        str(len(core)),
        str(len(recognition)),
        str(len(excluded)),
        pace,
    ]
    for index, value in enumerate(headers):
        base.set_cell_shading(summary_table.cell(0, index), base.LIGHT_BLUE)
        base.set_cell_text(
            summary_table.cell(0, index),
            value,
            bold=True,
            color=base.DARK_BLUE,
            align="center",
        )
        base.set_cell_text(
            summary_table.cell(1, index),
            values[index],
            size=11,
            bold=True,
            color=base.INK,
            align="center",
        )

    base.add_lead_callout(
        doc,
        "学习目标",
        f"本单元保留{len(retained)}条学习词目。先掌握A类核心，再完成B类认读，"
        "最后用核心词自测检查拼写。",
    )
    if excluded:
        base.add_lead_callout(
            doc,
            "本单元筛选",
            f"已隐藏{len(excluded)}条不要求背诵的专名或低考试价值文化专名。",
            fill=base.LIGHT_GRAY,
            accent=base.MUTED,
        )

    priority_rows = sorted(
        core,
        key=lambda row: (rank_value(row), int(row["source_row"])),
    )[:6]
    if priority_rows:
        words = "、".join(row["word"] for row in priority_rows)
        base.add_lead_callout(
            doc,
            "优先记忆",
            f"建议先掌握：{words}。",
            fill=base.GOLD_FILL,
            accent=base.GOLD,
        )


def select_test_rows(rows, limit=6):
    core = sorted(
        (row for row in rows if row["classification"] == "A核心"),
        key=lambda row: (rank_value(row), int(row["source_row"])),
    )
    selected = []
    seen = set()
    for row in core:
        key = clean_text(row["normalized_word"]).lower()
        if not key or key in seen:
            continue
        seen.add(key)
        selected.append(row)
        if len(selected) == limit:
            break
    return selected


def add_unit_test(doc, rows):
    base.add_heading(doc, "核心词自测", 2)
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(6)
    base.add_text(
        intro,
        "根据词性和教材释义写出单词或短语，再到书末核对答案。",
        size=9.5,
        color=base.MUTED,
    )

    selected = select_test_rows(rows)
    table = doc.add_table(rows=1, cols=2)
    widths = [700, 8660]
    base.set_table_geometry(table, widths)
    set_header_row(table, ["序号", "拼写自测"], fill=base.DARK_BLUE)

    for index, row in enumerate(selected, start=1):
        cells = table.add_row().cells
        fill = base.LIGHT_GRAY if index % 2 else base.WHITE
        for cell in cells:
            base.set_cell_shading(cell, fill)
        prompt_parts = [
            clean_text(row["part_of_speech"]),
            clean_text(row["chinese_definition"]),
        ]
        prompt = "  ".join(part for part in prompt_parts if part)
        base.set_cell_text(
            cells[0],
            str(index),
            bold=True,
            color=base.BLUE,
            align="center",
        )
        base.set_cell_text(
            cells[1],
            f"________________  {prompt}",
            size=9.8,
        )
    base.set_table_geometry(table, widths)

    base.add_heading(doc, "复习记录", 3)
    tracker = doc.add_table(rows=2, cols=7)
    tracker_widths = [900, 1410, 1410, 1410, 1410, 1410, 1410]
    base.set_table_geometry(tracker, tracker_widths)
    labels = ["复习", "第1天", "第2天", "第4天", "第7天", "第14天", "第30天"]
    for index, label in enumerate(labels):
        base.set_cell_shading(tracker.cell(0, index), base.LIGHT_BLUE)
        base.set_cell_text(
            tracker.cell(0, index),
            label,
            bold=True,
            color=base.DARK_BLUE,
            align="center",
        )
        base.set_cell_text(
            tracker.cell(1, index),
            "完成" if index == 0 else "□",
            size=11,
            align="center",
        )
    return [row["word"] for row in selected]


def add_unit(doc, book, unit, rows, page_break_before):
    add_unit_overview(doc, book, unit, rows, page_break_before)
    core = [row for row in rows if row["classification"] == "A核心"]
    recognition = [row for row in rows if row["classification"] == "B认读"]
    base.add_word_table(doc, "A类核心词", core, "A核心")
    base.add_word_table(doc, "B类认读词", recognition, "B认读")
    return add_unit_test(doc, rows)


def add_answer_key(doc, grouped, answers):
    doc.add_page_break()
    heading = base.add_heading(doc, "参考答案", 1)
    heading.paragraph_format.space_before = Pt(0)
    base.add_lead_callout(
        doc,
        "使用方式",
        "建议先独立完成每个单元的核心词自测，再按册次和单元核对答案。",
    )

    for book_index, book in enumerate(BOOK_ORDER):
        book_heading = base.add_heading(
            doc,
            f"{BOOK_NAMES[book]} {book}",
            2,
        )
        book_heading.paragraph_format.page_break_before = book_index > 0

        table = doc.add_table(rows=1, cols=2)
        widths = [2100, 7260]
        base.set_table_geometry(table, widths)
        set_header_row(table, ["单元", "核心词自测答案"], fill=base.DARK_BLUE)

        for row_index, unit in enumerate(grouped[book], start=1):
            cells = table.add_row().cells
            fill = base.WHITE if row_index % 2 else base.LIGHT_GRAY
            for cell in cells:
                base.set_cell_shading(cell, fill)
            answer_text = "；".join(
                f"{index}. {word}"
                for index, word in enumerate(answers[(book, unit)], start=1)
            )
            base.set_cell_text(
                cells[0],
                unit,
                size=9.5,
                bold=True,
                color=base.DARK_BLUE,
                align="center",
            )
            base.set_cell_text(cells[1], answer_text, size=9.5)
        base.set_table_geometry(table, widths)

    base.add_lead_callout(
        doc,
        "学习提醒",
        "答对后仍建议按第1、2、4、7、14、30天的节奏复习，直到能够快速准确拼写。",
        fill=base.GOLD_FILL,
        accent=base.GOLD,
    )


def build_document(output_path):
    rows = read_rows()
    grouped = group_rows(rows)
    validate_rows(grouped)

    doc = Document()
    configure_full_document(doc)
    doc.core_properties.title = "新教材同步词汇手册完整本"
    doc.core_properties.subject = "7A至9B教材同步词汇"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""

    add_cover(doc)
    add_how_to_use(doc)
    add_front_matter(doc, grouped)

    answers = {}
    generated_entries = 0
    for book in BOOK_ORDER:
        add_book_divider(doc, book, grouped[book])
        for unit_index, (unit, unit_rows) in enumerate(grouped[book].items()):
            retained = [
                row
                for row in unit_rows
                if row["classification"] in {"A核心", "B认读"}
            ]
            generated_entries += len(retained)
            answers[(book, unit)] = add_unit(
                doc,
                book,
                unit,
                unit_rows,
                page_break_before=unit_index > 0,
            )

    if generated_entries != EXPECTED_RETAINED:
        raise RuntimeError(f"Generated entry count mismatch: {generated_entries}")

    add_answer_key(doc, grouped, answers)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)

    print(f"output={output_path}")
    print(f"books={len(grouped)}")
    print(f"units={sum(len(units) for units in grouped.values())}")
    print(f"entries={generated_entries}")


def main():
    args = parse_args()
    build_document(args.output)


if __name__ == "__main__":
    main()
